"""Actual media replacement, equivalent URLs, pending races and scoped PPT errors."""
import os
import re
os.environ['MINDEX_SMOKE_OFFLINE'] = '1'
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit, urljoin
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from smoke_app import start_local_app_server, sync_playwright, launch_chromium


def deck_bytes(width, height):
    deck = Presentation()
    deck.slide_width, deck.slide_height = width * 100000, height * 100000
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, deck.slide_height)
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(20, 160, 80)
    shape.line.fill.background()
    output = BytesIO(); deck.save(output)
    return output.getvalue()


server, url = start_local_app_server()
try:
    with sync_playwright() as p:
        for engine in ([os.environ['MINDEX_SMOKE_ENGINE']] if os.environ.get('MINDEX_SMOKE_ENGINE') else ('chromium', 'webkit')):
            browser = launch_chromium(p) if engine == 'chromium' else p.webkit.launch()
            page = browser.new_page(viewport={'width': 1200, 'height': 900})
            held = []
            def route(r):
                name = urlsplit(r.request.url).path
                if urlsplit(r.request.url).netloc != urlsplit(url).netloc:
                    return r.abort()
                if name == '/identity-slow.svg':
                    held.append(r); return
                if name.startswith('/identity-image-'):
                    w, h = (90, 160) if '-b.' in name else (160, 90)
                    return r.fulfill(body=f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}"><rect width="{w}" height="{h}" fill="green"/></svg>', content_type='image/svg+xml')
                if name.startswith('/tests/fixtures/media-fit-'):
                    data=(Path(__file__).parent/'fixtures'/name.rsplit('/',1)[1]).read_bytes()
                    headers={'Accept-Ranges':'bytes'};status=200
                    if r.request.headers.get('range'):
                        match=re.fullmatch(r'bytes=(\d+)-(\d*)',r.request.headers['range']);assert match
                        start=int(match[1]);end=min(int(match[2]) if match[2] else len(data)-1,len(data)-1)
                        headers['Content-Range']=f'bytes {start}-{end}/{len(data)}';data=data[start:end+1];status=206
                    headers['Content-Length']=str(len(data))
                    return r.fulfill(status=status,body=data,headers=headers,content_type='video/mp4')
                if name.startswith('/identity-deck-'):
                    if 'missing' in name: return r.fulfill(status=404, body='Missing fixture')
                    return r.fulfill(body=deck_bytes(*( (90,160) if '-b.' in name else (160,90) )), content_type='application/octet-stream')
                return r.continue_()
            page.route('**/*', route)
            page.goto(url + '?mindexSmokeRaw=1')
            page.wait_for_function('typeof state!=="undefined"&&!document.body.classList.contains("ui-booting")')
            page.evaluate('''()=>{
              document.body.className='presenter-output-body';document.body.innerHTML='<main id="presenterOutputRoot" class="presenter-output-root"></main>';
              window.checkIdentity=(ok,message)=>{if(!ok)throw Error(message)};
              window.identityRoot=document.querySelector('main');
              checkIdentity(presenterMediaSourceIdentity('/a/../movie.mp4')===presenterMediaSourceIdentity(new URL('/movie.mp4',location.href).href),'equivalent URL normalization changed identity');
              checkIdentity(presenterMediaSourceIdentity('/movie.mp4?v=1')!==presenterMediaSourceIdentity('/movie.mp4?v=2')&&presenterMediaSourceIdentity('/movie.mp4#t=1')!==presenterMediaSourceIdentity('/movie.mp4#t=2'),'meaningful query/fragment discarded');
              checkIdentity(!presenterMediaSourceIdentity('javascript:alert(1)')&&!presenterMediaSourceIdentity('data:video/mp4,test'),'unsafe source became valid');
              window.identityPayload=(kind,source)=>({serviceId:'same-service',index:0,chromakey:false,messageId:'same-message',slides:[{
                id:'same-item',elementId:'same-item',type:kind,elementType:kind,layout:'media',title:'Same title',text:'Same title',
                ...(kind==='image'?{imageSrc:source}:kind==='video'?{videoSrc:source,playback:{autoplay:false,muted:true,loop:true,volume:.4}}:{pptx:{url:source,page:0,clicks:0}})
              }]});
              window.identityInputs=[];
              window.renderIdentity=(kind,source)=>{window.currentIdentity=identityPayload(kind,source);if(kind!=='pptx')identityInputs.push({input:currentIdentity,snapshot:JSON.stringify(currentIdentity)});renderPresenterOutput(currentIdentity)};
            }''')
            for kind, sources in [('image',['/identity-image-a.svg','/identity-image-b.svg']),
                                  ('video',['/tests/fixtures/media-fit-landscape.mp4','/tests/fixtures/media-fit-portrait.mp4'])]:
                print('CHECK',engine,kind,'replacement and same-source playback',flush=True)
                selector = '.is-active ' + ('img' if kind == 'image' else 'video')
                page.evaluate('args=>renderIdentity(...args)', [kind,sources[0]])
                page.wait_for_function('args=>document.querySelector(args[0])?.getAttribute("src")===args[1]', arg=[selector,sources[0]])
                page.wait_for_function('s=>{const e=document.querySelector(s);return e.tagName==="IMG"?e.naturalWidth>0:e.readyState>=2}', arg=selector)
                page.evaluate('''async s=>{
                  window.identityNode=document.querySelector(s);
                  if(identityNode.tagName==='VIDEO'){identityNode.volume=.4;await identityNode.play();await new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)));identityNode.pause();}
                  window.identityTime=identityNode.currentTime;
                }''', selector)
                for source in [sources[0], urljoin(url, sources[0])]:
                    page.evaluate('args=>renderIdentity(...args)', [kind,source])
                    page.evaluate('''s=>{
                      checkIdentity(document.querySelector(s)===identityNode,'same/equivalent source recreated media');
                      if(identityNode.tagName==='VIDEO')checkIdentity(identityNode.paused&&Math.abs(identityNode.currentTime-identityTime)<.03&&identityNode.loop&&identityNode.muted&&Math.abs(identityNode.volume-.4)<.01,'same source reset playback');
                    }''', selector)
                if kind == 'video':
                    page.evaluate('''()=>{
                      const health=presenterOutputVideoHealth(currentIdentity);
                      state.presenter.serviceId=currentIdentity.serviceId;state.presenter.index=0;state.presenter.slides=currentIdentity.slides;
                      state.presenter.safetyBlank=false;state.presenter.liveScripture=null;
                      state.presenter.videoHealth={...health,receivedAt:Date.now()};presenterTransportLatest={payload:{messageId:health.messageId}};
                      checkIdentity(presenterRemoteVideoState(currentIdentity.serviceId).ready,'matching source disabled');
                      state.presenter.slides=identityPayload('video','/different.mp4').slides;
                      checkIdentity(!presenterRemoteVideoState(currentIdentity.serviceId).ready,'old frame controls a new source');
                      checkIdentity(presenterVideoSourceFeedback(presenterRemoteVideoState(currentIdentity.serviceId))==='새 영상 반영 대기','source mismatch guidance missing');
                      checkIdentity(health.source===presenterMediaSourceIdentity(identityNode.currentSrc),'observed source not bound to real video');
                    }''')
                page.evaluate('args=>renderIdentity(...args)', [kind,sources[1]])
                page.wait_for_function('args=>document.querySelector(args[0])?.getAttribute("src")===args[1]', arg=[selector,sources[1]])
                # WebKit may expose a CSS-sized SVG raster as its natural dimensions.
                page.wait_for_function('s=>{const e=document.querySelector(s);return e.tagName==="IMG"?Math.abs(e.naturalWidth/e.naturalHeight-90/160)<.002:e.videoWidth===90&&e.readyState>=2}', arg=selector)
                assert page.eval_on_selector(selector, 'e=>e!==identityNode')
            # A pending image must not replace a newer committed source when it finishes late.
            page.evaluate("renderIdentity('image','/identity-slow.svg')")
            page.wait_for_function('presenterOutputImagePreloadCache.has("/identity-slow.svg")')
            page.evaluate("renderIdentity('image','/identity-image-c.svg')")
            page.wait_for_function('document.querySelector(".is-active img")?.getAttribute("src")==="/identity-image-c.svg"')
            assert held, 'Delayed source request was not exercised'
            for request in held:
                request.fulfill(body='<svg xmlns="http://www.w3.org/2000/svg" width="80" height="80"></svg>', content_type='image/svg+xml')
            page.wait_for_function('presenterOutputImagePreloadCache.get("/identity-slow.svg")?.image?.complete')
            page.evaluate('''async()=>{await new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)));checkIdentity(document.querySelector('.is-active img')?.getAttribute('src')==='/identity-image-c.svg','late source won crossfade race')}''')
            # PPTX already owns a deck source key; equivalent URLs must keep its native session.
            for source in ['/identity-deck-a.pptx', urljoin(url,'/identity-deck-a.pptx'), '/identity-deck-b.pptx']:
                page.evaluate('source=>renderIdentity("pptx",source)', source)
                page.wait_for_function('identityRoot._pptxMetadata?.key===presenterPptxKey(currentIdentity.slides[0])&&identityRoot.querySelector(".is-active .presenter-pptx-stage")?.dataset.pptxStep==="0"')
                page.evaluate('''source=>{
                  const svg=identityRoot.querySelector('.is-active svg'),session=identityRoot.querySelector('.is-active')._pptxSession;
                  if(source.startsWith('http'))checkIdentity(svg===window.identitySvg&&session===window.identitySession,'equivalent PPTX source reset session');
                  const box=svg.viewBox.baseVal,wanted=source.includes('-b.')?90/160:160/90;
                  checkIdentity(Math.abs(box.width/box.height-wanted)<.001,'PPTX actual source did not change');
                  window.identitySvg=svg;window.identitySession=session;
                }''', source)
            page.evaluate("renderIdentity('pptx','/identity-deck-missing.pptx')")
            page.wait_for_function('Boolean(identityRoot.dataset.pptxError)')
            page.evaluate('''()=>{checkIdentity(identityRoot.querySelector('.is-active svg')===identitySvg,'failed PPTX destroyed old output');checkIdentity(Boolean(presenterPptxOutputError(identityRoot,currentIdentity)),'current PPTX load error lost')}''')
            page.evaluate("renderIdentity('pptx','/identity-deck-b.pptx')")
            page.wait_for_function('!identityRoot.dataset.pptxError')
            page.evaluate('''()=>{
              const scope=presenterPptxScope(currentIdentity,currentIdentity.slides[0]);
              checkIdentity(identityRoot.querySelector('.is-active svg')===identitySvg,'recovery recreated retained PPTX session');
              identityRoot.dataset.pptxError='Other source error';identityRoot.dataset.pptxErrorScope='other-source';
              checkIdentity(!presenterPptxOutputError(identityRoot,currentIdentity),'wrong source error leaked');
              setPresenterPptxError(identityRoot,'Current source error',scope);setPresenterPptxError(identityRoot,'','other-source');
              checkIdentity(presenterPptxOutputError(identityRoot,currentIdentity)==='Current source error','old source cleared current error');setPresenterPptxError(identityRoot,'',scope);
              checkIdentity(identityInputs.every(({input,snapshot})=>JSON.stringify(input)===snapshot),'media source/playback settings mutated');
              releasePresenterPptxOutput(identityRoot);identityRoot.dataset.pptxError='Other active media error';
              setPresenterPptxError(identityRoot,'Late old PPT error',scope);setPresenterPptxError(identityRoot,'',scope);
              checkIdentity(presenterPptxOutputError(identityRoot,{slides:[],index:0})==='Other active media error','released PPT affected another media error');
            }''')
            print('PASS',engine,'image/video source replacement; URL-equivalent playback/session retention; observed-source control binding; late crossfade; actual PPTX source and scoped error recovery',flush=True)
            browser.close()
finally:
    server.shutdown()
