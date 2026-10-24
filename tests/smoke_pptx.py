import os
os.environ["MINDEX_SMOKE_OFFLINE"] = "1"
"""PPTX import/navigation/output, isolated storage and atomic-save doubles only."""
from io import BytesIO
import zipfile
from urllib.parse import urlsplit
from pptx import Presentation
from pptx.util import Inches
from lxml import etree
from PIL import Image
from smoke_app import launch_chromium, start_local_app_server, sync_playwright


def fixture():
    deck = Presentation()
    for name in ("Click 1", "Click 2", "Automatic", "Last slide"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        box.text = name
    picture = BytesIO(); Image.new("RGB", (80, 40), (30, 70, 120)).save(picture, format="PNG"); picture.seek(0)
    deck.slides[0].shapes.add_picture(picture, Inches(1), Inches(4))
    source = BytesIO(); deck.save(source)
    result = BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename in ("ppt/slides/slide1.xml", "ppt/slides/slide3.xml"):
                auto = info.filename.endswith("slide3.xml")
                condition = '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>' if auto else '<p:cond delay="indefinite"/>'
                timing = f'''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" nodeType="tmRoot"><p:childTnLst><p:seq><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst>{condition}</p:stCondLst><p:childTnLst><p:par><p:cTn id="4" presetID="1" presetClass="entr" nodeType="clickEffect" fill="hold"><p:childTnLst><p:set><p:cBhvr><p:cTn id="5" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="2"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
                xml = etree.fromstring(data); xml.append(etree.fromstring(timing)); data = etree.tostring(xml)
            output.writestr(info.filename, data)
    return result.getvalue()


def main():
    server, url = start_local_app_server()
    deck = fixture()
    try:
        with sync_playwright() as p:
            for engine in ("chromium", "webkit"):
                browser = launch_chromium(p) if engine == "chromium" else p.webkit.launch()
                page = browser.new_page(viewport={"width": 1280, "height": 720})
                page.route("**/*", lambda r: r.continue_() if (urlsplit(r.request.url).hostname in ("127.0.0.1", "localhost") or r.request.url.startswith(("blob:", "data:"))) else r.abort())
                page.route("**/test-deck.pptx", lambda r: r.fulfill(body=deck, content_type="application/octet-stream"))
                page.goto(url, wait_until="domcontentloaded")
                page.wait_for_function("typeof state !== 'undefined' && typeof importPresenterPptx === 'function' && !document.body.classList.contains('ui-booting')")
                result = page.evaluate("""async () => {
                  const check=(v,m)=>{if(!v)throw Error(m)};
                  const source=new URL('/test-deck.pptx',location.href).href;
                  const bytes=await(await fetch(source)).arrayBuffer();
                  const file=new File([bytes],'test.pptx',{type:'application/vnd.openxmlformats-officedocument.presentationml.presentation'});
                  const service={id:'ppt-fixture',type_id:'fixture'};
                  const item=normalizeServiceItem({id:'ppt-item',service_id:service.id,label:'참고 화면',_worshipSectionKey:'sermon',
                    memo:serializeServiceItemMemo({elementType:'image',inputMode:'asset',asset:{kind:'image'},note:'Preserve me'})},0);
                  state.services=[service];state.selectedServiceId=service.id;state.serviceItems={[service.id]:[item]};
                  window.uploads=[];window.removals=[];window.saved=0;window.expectedIndex=0;window.failUpload=false;window.failSave=false;
                  state.client={storage:{from:()=>({upload:async(path,blob)=>{
                    if(failUpload && uploads.length%2) return {error:new Error('upload failure')};
                    uploads.push({path,size:blob.size,type:blob.type});return {error:null};},
                    getPublicUrl:path=>({data:{publicUrl:path.endsWith('.pptx')?source:new URL('/'+path,location.href).href}}),
                    remove:async paths=>{removals.push(...paths);return {error:null}}})}};
                  window.confirm=()=>true;showToast=()=>{};renderCurrentServiceModuleDetail=()=>{};renderServiceList=()=>{};updateSaveState=()=>{};
                  markServiceElementDirty=()=>{};
                  saveServiceItemMutation=async(sid,index,options)=>{check(sid===service.id&&index===expectedIndex&&options.throwOnError,'wrong save boundary');saved++;if(failSave)throw Error('uncertain save');return true};
                  const input={files:[file],dataset:{serviceId:service.id},disabled:false,value:'x'};
                  check(await importPresenterPptx(input,item),'import failed');
                  check(uploads.length===5&&saved===1&&!input.disabled&&!input.value,'import lifecycle');
                  const memo=parseServiceItemMemo(item.memo),asset=normalizeServiceAsset(memo.asset);
                  check(memo.note==='Preserve me','metadata overwritten');
                  check(asset.kind==='imported_deck'&&asset.pptx.url===source&&asset.slides.length===4,'asset lost');
                  check(asset.original.name===file.name&&asset.original.format==='pptx'&&asset.original.url===source&&asset.original.sha256===await presenterPptxDigest(await file.arrayBuffer()),'PPTX original metadata lost');
                  check(asset.slides[0].pptx.clicks===1&&asset.slides[2].pptx.clicks===0,'click/auto counts');
                  const slides=presenterImportedDeckSlidesFromAsset(asset,item,{sectionKey:'sermon'},0,'PPT','자료');
                  check(slides.length===4&&slides[0].pptx.url===source,'presenter build');
                  const persisted=normalizeServiceDocumentSlides(slides);
                  check(persisted[0].pptx.page===0&&persisted[0].asset.pptx.url===source&&persisted[0].asset.original.sha256===asset.original.sha256,'document round trip');
                  window.pp={serviceId:service.id,index:0,slides};
                  let moved=applyPresenterActionToPayload(pp,'next');
                  check(moved.index===0&&presenterPptxStep(moved)===1,'click skipped');
                  moved=applyPresenterActionToPayload(moved,'prev');check(moved.index===0&&presenterPptxStep(moved)===0,'rewind');
                  moved=applyPresenterActionToPayload(applyPresenterActionToPayload(moved,'next'),'next');check(moved.index===1,'did not advance');
                  moved=applyPresenterActionToPayload({...pp,pptxProgress:{key:presenterPptxKey(slides[0]),step:1}},'jump',{index:0});check(presenterPptxStep(moved)===0,'jump did not reset');
                  check(applyPresenterActionToPayload(pp,'jump',{index:-1}).safetyBlank,'zero did not blank');
                  const clicked=applyPresenterActionToPayload(pp,'next');check(presenterPptxStep(applyPresenterActionToPayload(clicked,'jump',{index:999}))===1,'invalid jump reset click');
                  const before=item.memo;failUpload=true;uploads=[];removals=[];
                  check(!await importPresenterPptx(input,item),'upload failure accepted');
                  check(item.memo===before&&saved===1&&removals.length===1,'failed upload mutated item/leaked files');
                  failUpload=false;failSave=true;removals=[];
                  check(!await importPresenterPptx(input,item),'save failure accepted');
                  check(removals.length===0,'uncertain save deleted referenced media');
                  check(!input.disabled&&!presenterPptxImports.size,'busy guard stuck');
                  // Controller and output-only navigation use the same click semantics.
                  state.presenter.slides=slides;state.presenter.index=0;state.presenter.pptxProgress=null;
                  movePresenterSlide(1);check(state.presenter.index===0&&presenterPptxStep(state.presenter)===1,'controller skipped click');
                  movePresenterSlide(1);check(state.presenter.index===1,'controller stuck');
                  failSave=false;expectedIndex=1;
                  check(await importPresenterPptx({...input,dataset:{serviceId:service.id,presenterReferenceMediaSection:'sermon'}}),'quick add failed');
                  check(getServiceItems(service.id).length===2,'quick add did not insert exactly one item');
                  check(parseServiceItemMemo(getServiceItems(service.id)[1].memo).asset.pptx?.url===source,'quick add lost source');
                  return 'quick add, replacement, preservation, round trip, atomic boundary, failure retention, navigation';
                }""")
                print('PASS', engine, result, flush=True)
                page.evaluate("""() => {
                  document.body.innerHTML='<div id="presenterOutputRoot"></div>';
                  renderPresenterOutput({...pp,safetyBlank:true});
                  renderPresenterOutput(pp);
                }""")
                page.wait_for_function("document.querySelector('.presenter-pptx-stage')?.dataset.pptxStep==='0'")
                page.evaluate("""() => {
                  const check=(v,m)=>{if(!v)throw Error(m)};
                  const target=document.querySelector('.presenter-pptx-stage [data-el="2"]');
                  check(getComputedStyle(target).visibility==='hidden','entrance visible before click');
                  pp=applyPresenterActionToPayload(pp,'next');renderPresenterOutput(pp);
                }""")
                page.wait_for_timeout(80)
                page.evaluate("""() => {
                  const target=document.querySelector('.presenter-pptx-stage [data-el="2"]');
                  if(getComputedStyle(target).visibility==='hidden'||Number(getComputedStyle(target).opacity)!==1) throw Error('instant appearance failed');
                  window.sameHost=document.querySelector('.presenter-pptx-stage');renderPresenterOutput(pp);
                  if(sameHost!==document.querySelector('.presenter-pptx-stage'))throw Error('duplicate replay');
                  pp=applyPresenterActionToPayload(pp,'prev');renderPresenterOutput(pp);
                  if(getComputedStyle(document.querySelector('.presenter-pptx-stage [data-el="2"]')).visibility!=='hidden')throw Error('rewind visible');
                  pp=applyPresenterActionToPayload(pp,'jump',{index:2});renderPresenterOutput(pp);
                }""")
                page.wait_for_function("document.querySelector('.presenter-output-layer.is-active .presenter-pptx-stage')?.textContent.includes('Automatic')")
                page.wait_for_timeout(80)
                page.evaluate("""() => {
                  const target=document.querySelector('.presenter-output-layer.is-active [data-el="2"]');
                  if(getComputedStyle(target).visibility==='hidden'||Number(getComputedStyle(target).opacity)!==1)throw Error('automatic start failed');
                  renderPresenterOutput({...pp,safetyBlank:true});
                  if(!document.getElementById('presenterOutputRoot').classList.contains('is-blackout'))throw Error('blackout failed');
                  if(document.querySelector('.presenter-pptx-stage'))throw Error('PPT survived blackout');
                }""")
                pending=[]
                page.route('**/slow-deck.pptx', lambda r: pending.append(r))
                page.evaluate("""() => {
                  const slow={...pp,index:0,slides:[{...pp.slides[0],pptx:{...pp.slides[0].pptx,url:new URL('/slow-deck.pptx',location.href).href}}]};
                  renderPresenterOutput(slow);
                }""")
                page.wait_for_timeout(100)
                assert pending and page.evaluate("document.getElementById('presenterOutputRoot').classList.contains('is-blackout')"), 'loading replaced live frame'
                page.evaluate("renderPresenterOutput({...pp,safetyBlank:true})")
                for route in pending: route.fulfill(body=deck,content_type='application/octet-stream')
                page.wait_for_timeout(200)
                assert page.evaluate("document.getElementById('presenterOutputRoot').classList.contains('is-blackout') && !document.querySelector('.presenter-pptx-stage')"), 'stale load overrode blackout'
                print('PASS',engine,'real output: appear, auto, rewind, replay, blackout, delayed load',flush=True)
                browser.close()
    finally:
        server.shutdown();server.server_close()


if __name__ == '__main__':
    main()
